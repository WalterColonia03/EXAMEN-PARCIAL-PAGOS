"""Small, deterministic import helpers for the Stage 1 service-account registry."""
from __future__ import annotations

import csv
import io
import unicodedata
from collections.abc import Iterable
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile


class ServiceAccountImportError(ValueError):
    """A friendly, row-aware error that callers can display directly."""


REQUIRED_COLUMNS = {
    "creditor": "acreedor / proveedor",
    "service_type": "tipo de servicio",
    "identifier_type": "tipo de identificador",
    "primary_identifier": "identificador principal",
}

HEADER_ALIASES = {
    "ambito": "context",
    "contexto": "context",
    "context": "context",
    "empresa": "company",
    "compania": "company",
    "company": "company",
    "local": "branch",
    "sucursal": "branch",
    "branch": "branch",
    "acreedor": "creditor",
    "proveedor": "creditor",
    "creditor": "creditor",
    "tipo servicio": "service_type",
    "tipo de servicio": "service_type",
    "servicio": "service_type",
    "service type": "service_type",
    "tipo identificador": "identifier_type",
    "tipo de identificador": "identifier_type",
    "identifier type": "identifier_type",
    "identificador principal": "primary_identifier",
    "codigo principal": "primary_identifier",
    "codigo": "primary_identifier",
    "suministro": "primary_identifier",
    "primary identifier": "primary_identifier",
    "titular": "account_holder",
    "account holder": "account_holder",
    "frecuencia": "expected_billing_frequency",
    "frecuencia facturacion": "expected_billing_frequency",
    "billing frequency": "expected_billing_frequency",
    "dia habitual": "usual_billing_day",
    "dia de facturacion": "usual_billing_day",
    "billing day": "usual_billing_day",
    "canal": "acquisition_channel",
    "canal de obtencion": "acquisition_channel",
    "acquisition channel": "acquisition_channel",
    "estado": "status",
    "status": "status",
    "identificadores adicionales": "additional_identifiers",
    "additional identifiers": "additional_identifiers",
}

SERVICE_TYPES = {
    "electricidad": "ELECTRICITY",
    "electrico": "ELECTRICITY",
    "electricity": "ELECTRICITY",
    "agua": "WATER",
    "water": "WATER",
    "gas": "GAS",
    "internet": "INTERNET",
    "telecomunicaciones": "TELECOM",
    "telecom": "TELECOM",
    "municipal": "MUNICIPAL",
    "arbitrios": "MUNICIPAL",
    "tributo": "TAX",
    "impuesto": "TAX",
    "tax": "TAX",
    "otro": "OTHER",
    "other": "OTHER",
}

IDENTIFIER_TYPES = {
    "suministro": "SUPPLY",
    "supply": "SUPPLY",
    "cliente": "CLIENT",
    "client": "CLIENT",
    "contrato": "CONTRACT",
    "contract": "CONTRACT",
    "predio": "PROPERTY",
    "property": "PROPERTY",
    "contribuyente": "TAXPAYER",
    "taxpayer": "TAXPAYER",
    "otro": "OTHER",
    "other": "OTHER",
}

FREQUENCIES = {
    "mensual": "MONTHLY",
    "monthly": "MONTHLY",
    "bimestral": "BIMONTHLY",
    "bimonthly": "BIMONTHLY",
    "trimestral": "QUARTERLY",
    "quarterly": "QUARTERLY",
    "eventual": "ON_DEMAND",
    "segun consumo": "ON_DEMAND",
    "on demand": "ON_DEMAND",
}

ACQUISITION_CHANNELS = {
    "correo": "EMAIL",
    "email": "EMAIL",
    "portal": "PORTAL",
    "api": "API",
    "ocr": "OCR",
    "documento digitalizado": "OCR",
    "manual": "MANUAL",
}

STATUSES = {
    "activa": "ACTIVE",
    "activo": "ACTIVE",
    "active": "ACTIVE",
    "suspendida": "SUSPENDED",
    "suspendido": "SUSPENDED",
    "suspended": "SUSPENDED",
    "cerrada": "CLOSED",
    "cerrado": "CLOSED",
    "closed": "CLOSED",
}


def _normalized(value: object | None) -> str:
    raw = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", raw)
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(without_marks.strip().lower().replace("_", " ").split())


def _text(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalise_headers(headers: Iterable[object | None]) -> list[str | None]:
    normalised = []
    for header in headers:
        normalised.append(HEADER_ALIASES.get(_normalized(header)))
    return normalised


def _parse_additional_identifiers(value: str, row_number: int) -> dict[str, str]:
    if not value:
        return {}
    identifiers: dict[str, str] = {}
    for pair in value.split("|"):
        label, separator, identifier = pair.partition(":")
        label, identifier = label.strip(), identifier.strip()
        if not separator or not label or not identifier:
            raise ServiceAccountImportError(
                f"Fila {row_number}: los identificadores adicionales deben usar Nombre: valor, separados por |."
            )
        if label in identifiers:
            raise ServiceAccountImportError(f"Fila {row_number}: el identificador adicional «{label}» se repite.")
        identifiers[label] = identifier
    return identifiers


def _choice(value: str, mapping: dict[str, str], field_label: str, row_number: int, default: str | None = None) -> str:
    if not value and default is not None:
        return default
    selected = mapping.get(_normalized(value))
    if selected is None:
        raise ServiceAccountImportError(
            f"Fila {row_number}: «{value}» no es válido para {field_label}. "
            "Descargue la plantilla y use una opción admitida."
        )
    return selected


def _rows_from_csv(upload) -> list[list[object | None]]:
    try:
        content = upload.read()
        decoded = content.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(decoded), strict=True))
    except UnicodeDecodeError as error:
        raise ServiceAccountImportError("El CSV debe estar codificado en UTF-8.") from error
    except csv.Error as error:
        raise ServiceAccountImportError("El CSV está dañado o contiene filas con formato inválido.") from error
    finally:
        upload.seek(0)


def _rows_from_xlsx(upload) -> list[list[object | None]]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:  # pragma: no cover - requirements pin this package.
        raise ServiceAccountImportError(
            "No fue posible leer el archivo Excel. Inténtalo nuevamente o utiliza la plantilla CSV."
        ) from error

    workbook = None
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
        IndexError,
        ParseError,
    ) as error:
        raise ServiceAccountImportError("El archivo Excel está dañado o no es un libro .xlsx válido.") from error
    finally:
        if workbook is not None:
            workbook.close()
        upload.seek(0)


def parse_service_account_import(upload) -> list[dict[str, object]]:
    """Parse a CSV/XLSX upload into validated, database-independent account rows."""
    suffix = Path(upload.name).suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(upload)
    elif suffix == ".xlsx":
        rows = _rows_from_xlsx(upload)
    else:
        raise ServiceAccountImportError("Seleccione un archivo CSV o Excel .xlsx.")

    if not rows or not any(_text(cell) for cell in rows[0]):
        raise ServiceAccountImportError("El archivo debe incluir una fila de encabezados.")

    headers = _normalise_headers(rows[0])
    missing = [label for key, label in REQUIRED_COLUMNS.items() if key not in headers]
    if "context" not in headers and "company" not in headers:
        missing.insert(0, "ámbito o empresa")
    if missing:
        raise ServiceAccountImportError("Faltan columnas obligatorias: " + ", ".join(missing) + ".")

    parsed: list[dict[str, object]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(_text(value) for value in values):
            continue
        source = {
            header: _text(values[index]) if index < len(values) else ""
            for index, header in enumerate(headers)
            if header
        }
        blank_fields = [label for key, label in REQUIRED_COLUMNS.items() if not source.get(key)]
        if not source.get("context") and not source.get("company"):
            blank_fields.insert(0, "ámbito o empresa")
        if blank_fields:
            raise ServiceAccountImportError(
                f"Fila {row_number}: faltan campos obligatorios: {', '.join(blank_fields)}."
            )
        billing_day = source.get("usual_billing_day", "")
        if billing_day:
            try:
                billing_day = int(billing_day)
            except (TypeError, ValueError) as error:
                raise ServiceAccountImportError(f"Fila {row_number}: el día habitual debe ser un número entre 1 y 31.") from error
            if not 1 <= billing_day <= 31:
                raise ServiceAccountImportError(f"Fila {row_number}: el día habitual debe estar entre 1 y 31.")
        else:
            billing_day = None

        parsed.append(
            {
                "row_number": row_number,
                "context_name": source.get("context", "") or source.get("company", ""),
                "company_name": source.get("company", ""),
                "branch_name": source.get("branch", ""),
                "creditor_name": source["creditor"],
                "service_type": _choice(source["service_type"], SERVICE_TYPES, "tipo de servicio", row_number),
                "identifier_type": _choice(
                    source["identifier_type"], IDENTIFIER_TYPES, "tipo de identificador", row_number
                ),
                "primary_identifier": source["primary_identifier"],
                "account_holder": source.get("account_holder", ""),
                "expected_billing_frequency": _choice(
                    source.get("expected_billing_frequency", ""),
                    FREQUENCIES,
                    "frecuencia",
                    row_number,
                    default="MONTHLY",
                ),
                "usual_billing_day": billing_day,
                "acquisition_channel": _choice(
                    source.get("acquisition_channel", ""),
                    ACQUISITION_CHANNELS,
                    "canal",
                    row_number,
                    default="MANUAL",
                ),
                "status": _choice(source.get("status", ""), STATUSES, "estado", row_number, default="ACTIVE"),
                "additional_identifiers": _parse_additional_identifiers(
                    source.get("additional_identifiers", ""), row_number
                ),
            }
        )

    if not parsed:
        raise ServiceAccountImportError("El archivo no contiene cuentas de servicio para importar.")
    return parsed
