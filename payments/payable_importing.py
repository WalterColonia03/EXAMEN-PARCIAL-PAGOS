"""Parse payable CSV/XLSX files for a review-first import flow."""
from __future__ import annotations

import csv
import io
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile


class PayableImportError(ValueError):
    pass


def _normalized(value) -> str:
    raw = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", raw)
    plain = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(plain.strip().lower().replace("_", " ").split())


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


HEADER_ALIASES = {
    "ambito": "context",
    "contexto": "context",
    "context": "context",
    "categoria": "category",
    "category": "category",
    "concepto": "concept",
    "concept": "concept",
    "acreedor": "creditor",
    "proveedor": "creditor",
    "creditor": "creditor",
    "documento": "external_document_number",
    "numero de documento": "external_document_number",
    "external document": "external_document_number",
    "periodo": "period",
    "period": "period",
    "monto": "amount",
    "amount": "amount",
    "moneda": "currency",
    "currency": "currency",
    "vencimiento": "due_date",
    "fecha de vencimiento": "due_date",
    "due date": "due_date",
    "titular": "holder_name",
    "deudor": "holder_name",
    "holder": "holder_name",
    "beneficiario": "beneficiary_name",
    "beneficia a": "beneficiary_name",
    "beneficiary": "beneficiary_name",
    "naturaleza": "payment_nature",
    "nature": "payment_nature",
    "sensible": "is_sensitive",
    "privado": "is_sensitive",
    "sensitive": "is_sensitive",
}

REQUIRED_COLUMNS = {
    "context": "ámbito",
    "category": "categoría",
    "concept": "concepto",
    "creditor": "acreedor",
    "due_date": "vencimiento",
}

CATEGORIES = {
    "servicio": "SERVICE",
    "vivienda": "HOUSING",
    "propiedad": "HOUSING",
    "tarjeta": "CREDIT_CARD",
    "tarjeta de credito": "CREDIT_CARD",
    "prestamo": "LOAN",
    "impuesto": "TAX",
    "seguro": "INSURANCE",
    "educacion": "EDUCATION",
    "salud": "HEALTH",
    "apoyo familiar": "FAMILY_SUPPORT",
    "personal domestico": "DOMESTIC_STAFF",
    "vehiculo": "VEHICLE",
    "viaje": "TRAVEL_EVENT",
    "evento": "TRAVEL_EVENT",
    "inversion": "INVESTMENT",
    "donacion": "DONATION",
    "legal": "LEGAL",
    "suscripcion": "SUBSCRIPTION",
    "proveedor": "SUPPLIER",
    "emergencia": "EMERGENCY",
    "otro": "OTHER",
}

NATURES = {
    "propio": "OWN",
    "compartido": "SHARED",
    "apoyo": "SUPPORT",
    "apoyo no reembolsable": "SUPPORT",
    "reembolsable": "REIMBURSABLE",
    "prestamo": "PERSONAL_LOAN",
    "prestamo a tercero": "PERSONAL_LOAN",
}

TRUE_VALUES = {"si", "sí", "true", "1", "x", "privado", "sensible"}
FALSE_VALUES = {"", "no", "false", "0"}


def _rows_from_csv(upload):
    try:
        return list(csv.reader(io.StringIO(upload.read().decode("utf-8-sig")), strict=True))
    except UnicodeDecodeError as error:
        raise PayableImportError("El CSV debe estar codificado en UTF-8.") from error
    except csv.Error as error:
        raise PayableImportError("El CSV está dañado o tiene filas inválidas.") from error
    finally:
        upload.seek(0)


def _rows_from_xlsx(upload):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:  # pragma: no cover
        raise PayableImportError("No fue posible leer el archivo Excel.") from error
    workbook = None
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        return [list(row) for row in workbook.active.iter_rows(values_only=True)]
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, ParseError) as error:
        raise PayableImportError("El archivo Excel está dañado o no es un .xlsx válido.") from error
    finally:
        if workbook is not None:
            workbook.close()
        upload.seek(0)


def _parse_date(value, label: str) -> str:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"{label} debe usar AAAA-MM-DD o DD/MM/AAAA")


def parse_payable_import(upload) -> list[dict]:
    suffix = Path(upload.name).suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(upload)
    elif suffix == ".xlsx":
        rows = _rows_from_xlsx(upload)
    else:
        raise PayableImportError("Seleccione un archivo CSV o Excel .xlsx.")
    if not rows or not any(_text(cell) for cell in rows[0]):
        raise PayableImportError("El archivo debe incluir encabezados.")
    headers = [HEADER_ALIASES.get(_normalized(cell)) for cell in rows[0]]
    missing = [label for key, label in REQUIRED_COLUMNS.items() if key not in headers]
    if missing:
        raise PayableImportError("Faltan columnas obligatorias: " + ", ".join(missing) + ".")

    parsed = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(_text(value) for value in values):
            continue
        source = {
            header: values[index] if index < len(values) else ""
            for index, header in enumerate(headers)
            if header
        }
        errors = []
        for key, label in REQUIRED_COLUMNS.items():
            if not _text(source.get(key)):
                errors.append(f"Falta {label}")
        category = CATEGORIES.get(_normalized(source.get("category")))
        if _text(source.get("category")) and category is None:
            errors.append("Categoría no reconocida")
        currency = _text(source.get("currency") or "PEN").upper()
        if currency not in {"PEN", "USD"}:
            errors.append("La moneda debe ser PEN o USD")
        amount = _text(source.get("amount"))
        if amount:
            try:
                decimal_amount = Decimal(amount.replace(",", "."))
                if decimal_amount <= 0:
                    raise InvalidOperation
                amount = str(decimal_amount.quantize(Decimal("0.01")))
            except (InvalidOperation, ValueError):
                errors.append("El monto debe ser un número mayor que cero")
                amount = ""
        try:
            due_date = _parse_date(source.get("due_date"), "El vencimiento")
        except ValueError as error:
            errors.append(str(error))
            due_date = ""
        period = ""
        if _text(source.get("period")):
            try:
                period = _parse_date(source.get("period"), "El período")
            except ValueError as error:
                errors.append(str(error))
        nature_raw = _normalized(source.get("payment_nature"))
        nature = NATURES.get(nature_raw, "OWN")
        if nature_raw and nature_raw not in NATURES:
            errors.append("Naturaleza no reconocida")
        sensitive_raw = _normalized(source.get("is_sensitive"))
        if sensitive_raw not in TRUE_VALUES | FALSE_VALUES:
            errors.append("Sensible debe indicar Sí o No")
        parsed.append(
            {
                "row_number": row_number,
                "context_name": _text(source.get("context")),
                "category": category or "OTHER",
                "concept": _text(source.get("concept")),
                "creditor_name": _text(source.get("creditor")),
                "external_document_number": _text(source.get("external_document_number")),
                "period": period,
                "amount": amount,
                "currency": currency,
                "due_date": due_date,
                "holder_name": _text(source.get("holder_name")),
                "beneficiary_name": _text(source.get("beneficiary_name")),
                "payment_nature": nature,
                "is_sensitive": sensitive_raw in TRUE_VALUES,
                "errors": errors,
            }
        )
    if not parsed:
        raise PayableImportError("El archivo no contiene obligaciones para revisar.")
    return parsed
